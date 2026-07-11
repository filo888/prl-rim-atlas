from __future__ import annotations

import json
import zipfile
from io import BytesIO

import numpy as np
import openpyxl
import pandas as pd
import pytest

from prl.analysis import analyze_case
from prl.demo import demo_files
from prl.dicom_io import decode_dicom, inspect_dicom
from prl.exports import export_csv, export_excel, export_json
from prl.ingest import UploadError, process_uploads


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return output.getvalue()


def test_synthetic_dicom_spacing_and_pixels_are_decoded(dicom_asset) -> None:
    image = np.arange(30, dtype=np.uint16).reshape(5, 6) * 100
    asset = dicom_asset(image, spacing=(0.7, 0.3))

    info = inspect_dicom(asset)
    decoded, decoded_info, warnings = decode_dicom(asset)

    assert (info.rows, info.cols, info.frames) == (5, 6, 1)
    assert info.spacing_mm == pytest.approx((0.7, 0.3))
    assert info.spacing_source == "DICOM PixelSpacing"
    assert decoded_info == info
    assert decoded.shape == image.shape
    assert decoded.dtype == np.uint8
    assert decoded[0, 0] < decoded[-1, -1]
    assert warnings == []


def test_only_the_selected_multiframe_slice_is_returned(dicom_asset) -> None:
    ascending = np.arange(20, dtype=np.uint16).reshape(4, 5)
    descending = ascending.max() - ascending
    asset = dicom_asset(np.stack([ascending, descending]))

    decoded, info, _ = decode_dicom(asset, frame_index=1)

    assert info.frames == 2
    assert decoded.shape == (4, 5)
    assert decoded[0, 0] > decoded[-1, -1]


def test_detector_spacing_is_not_silently_used_as_patient_spacing(
    dicom_asset,
) -> None:
    asset = dicom_asset(
        np.arange(16, dtype=np.uint16).reshape(4, 4),
        spacing=None,
        detector_spacing=(0.8, 0.6),
    )

    info = inspect_dicom(asset)

    assert info.spacing_mm is None
    assert info.spacing_source is None
    assert info.detector_spacing_mm == pytest.approx((0.8, 0.6))
    assert "detector plane" in info.detector_spacing_source


def test_process_uploads_reads_safe_zip_and_reports_ignored_files() -> None:
    payload = _zip_bytes(
        {
            "case/mask1.roi": b"Iout-test-roi",
            "case/image.dcm": b"test-dicom",
            "case/notes.txt": b"ignored",
        }
    )

    assets, notices = process_uploads([("batch.zip", payload)])

    assert [(asset.name, asset.kind) for asset in assets] == [
        ("case/mask1.roi", "roi"),
        ("case/image.dcm", "dicom"),
    ]
    assert all(asset.source_name == "batch.zip" for asset in assets)
    assert all(asset.archive_member is not None for asset in assets)
    assert any("ignored 1" in notice for notice in notices)


@pytest.mark.parametrize("unsafe_name", ["../escape.roi", "/absolute/image.dcm"])
def test_process_uploads_rejects_unsafe_zip_paths(unsafe_name: str) -> None:
    payload = _zip_bytes({unsafe_name: b"unsafe"})

    with pytest.raises(UploadError, match="unsafe ZIP member path"):
        process_uploads([("unsafe.zip", payload)])


def test_process_uploads_enforces_member_size_limit(monkeypatch) -> None:
    monkeypatch.setattr("prl.ingest.MAX_MEMBER_BYTES", 3)
    payload = _zip_bytes({"too-large.roi": b"1234"})

    with pytest.raises(UploadError, match="per-file limit"):
        process_uploads([("large.zip", payload)])


@pytest.fixture(scope="module")
def analyzed_demo():
    assets, notices = process_uploads(list(demo_files()))
    assert notices == []
    rim = next(asset for asset in assets if asset.name.endswith("mask1.roi"))
    core = next(asset for asset in assets if asset.name.endswith("mask2.roi"))
    dicom = next(asset for asset in assets if asset.kind == "dicom")
    return analyze_case(
        "DEMO_PRL1",
        rim_roi=rim,
        lesion_roi=core,
        dicom=dicom,
        spacing_policy="dicom_or_pixels",
        roi_mode="study_pixels",
        remove_overlap=True,
    )


def test_demo_files_run_end_to_end(analyzed_demo) -> None:
    result = analyzed_demo

    assert result.case_id == "DEMO_PRL1"
    assert result.canvas.shape == (160, 160)
    assert result.dicom_image is not None
    assert result.dicom_image.shape == result.canvas.shape
    assert result.rim_mask is not None and result.rim_mask.any()
    assert result.lesion_mask is not None and result.lesion_mask.any()
    assert not np.any(result.rim_mask & result.lesion_mask)
    assert result.spacing_mm == pytest.approx((0.55, 0.55))
    assert result.spacing_source == "DICOM PixelSpacing"
    assert result.metrics["mask1_area_rim_mm2"] == pytest.approx(
        result.metrics["mask1_area_rim_px"] * 0.55**2
    )
    assert result.metrics["rim_class_using_1p2_mm_threshold"] in {
        "BROAD RIM",
        "NARROW RIM",
    }


def test_all_exports_are_parseable_and_preserve_qc(analyzed_demo) -> None:
    result = analyzed_demo
    result.warnings.append("synthetic QC warning")
    errors = [{"case_name": "FAILED_CASE", "error": "synthetic failure"}]

    csv_bytes = export_csv([result])
    csv_frame = pd.read_csv(BytesIO(csv_bytes))
    assert csv_frame.loc[0, "case_name"] == "DEMO_PRL1"
    assert csv_frame.loc[0, "pixel_spacing_row_mm"] == pytest.approx(0.55)

    json_record = json.loads(export_json([result], errors).decode("utf-8"))
    assert json_record["cases"][0]["case_name"] == "DEMO_PRL1"
    assert "synthetic QC warning" in json_record["cases"][0]["warnings"]
    assert json_record["errors"] == errors

    workbook = openpyxl.load_workbook(BytesIO(export_excel([result], errors)))
    assert workbook.sheetnames == [
        "read_me",
        "lesion_metrics",
        "qc_provenance",
        "errors",
    ]
    assert workbook["lesion_metrics"].freeze_panes == "A2"
    assert workbook["qc_provenance"]["B2"].value == "warning"
    assert workbook["errors"]["A2"].value == "FAILED_CASE"


def test_spreadsheet_exports_neutralize_formula_like_user_text(analyzed_demo) -> None:
    analyzed_demo.metrics["case_name"] = "=HYPERLINK(\"https://example.invalid\")"
    analyzed_demo.case_id = "+SUM(1,1)"

    csv_text = export_csv([analyzed_demo]).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text

    workbook = openpyxl.load_workbook(BytesIO(export_excel([analyzed_demo])))
    metrics_sheet = workbook["lesion_metrics"]
    case_column = next(
        cell.column for cell in metrics_sheet[1] if cell.value == "case_name"
    )
    case_cell = metrics_sheet.cell(2, case_column)
    assert case_cell.data_type != "f"
    assert case_cell.value.startswith("'=")
