from __future__ import annotations

import json
import math
from io import BytesIO
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .models import CaseResult


HEADER_FILL = PatternFill("solid", fgColor="173844")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="E8F6F5")


def _clean_scalar(value: Any) -> Any:
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def _safe_spreadsheet_value(value: Any) -> Any:
    """Neutralize text that spreadsheet programs may interpret as a formula."""
    value = _clean_scalar(value)
    if not isinstance(value, str):
        return value
    stripped = value.lstrip()
    if value.startswith(("\t", "\r")) or stripped.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sanitize_frame(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.map(_safe_spreadsheet_value)


def results_dataframe(results: list[CaseResult]) -> pd.DataFrame:
    frame = pd.DataFrame(
        [
            {key: _clean_scalar(value) for key, value in result.metrics.items()}
            for result in results
        ]
    )
    return _sanitize_frame(frame)


def export_csv(results: list[CaseResult]) -> bytes:
    return results_dataframe(results).to_csv(index=False).encode("utf-8-sig")


def _style_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column[:250]) + 2
        sheet.column_dimensions[letter].width = min(max(width, 12), 48)


def export_excel(
    results: list[CaseResult], errors: list[dict[str, str]] | None = None
) -> bytes:
    metrics = results_dataframe(results)
    qc_rows = [
        {
            "case_name": result.case_id,
            "status": "warning" if result.warnings else "ready",
            "warning_count": len(result.warnings),
            "warnings": " | ".join(result.warnings),
            "spacing_source": result.spacing_source or "pixel metrics only",
            "dicom_file": result.dicom_file or "",
            "rim_file": result.rim_file or "",
            "lesion_file": result.lesion_file or "",
        }
        for result in results
    ]
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        metrics.to_excel(writer, index=False, sheet_name="lesion_metrics")
        _sanitize_frame(pd.DataFrame(qc_rows)).to_excel(
            writer, index=False, sheet_name="qc_provenance"
        )
        if errors:
            _sanitize_frame(pd.DataFrame(errors)).to_excel(
                writer, index=False, sheet_name="errors"
            )
        for sheet in writer.sheets.values():
            _style_sheet(sheet)
        summary = writer.book.create_sheet("read_me", 0)
        summary.append(["PRL RIM Atlas export"])
        summary.append(["Cases analyzed", len(results)])
        summary.append(["Cases failed", len(errors or [])])
        summary.append([
            "Method",
            "Paper-compatible point-pixel masks, 4-edge perimeters, deterministic medial-axis thickness.",
        ])
        summary.append([
            "Units",
            "Millimetre metrics are blank unless calibrated row/column spacing was available.",
        ])
        summary["A1"].fill = HEADER_FILL
        summary["A1"].font = HEADER_FONT
        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 92
        for row in range(2, 6):
            summary[f"A{row}"].fill = ACCENT_FILL
            summary[f"A{row}"].font = Font(bold=True, color="173844")
            summary[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    return buffer.getvalue()


def export_json(
    results: list[CaseResult], errors: list[dict[str, str]] | None = None
) -> bytes:
    payload = {
        "format": "PRL RIM Atlas reproducibility record",
        "cases": [
            {
                "case_name": result.case_id,
                "metrics": {key: _clean_scalar(value) for key, value in result.metrics.items()},
                "warnings": result.warnings,
            }
            for result in results
        ],
        "errors": errors or [],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")
